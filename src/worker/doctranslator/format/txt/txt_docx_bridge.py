"""Bridge plain-text (.txt) documents onto the native DOCX translation pipeline.

Rather than building a third format-specific translator, a `.txt` upload is
wrapped in a minimal in-memory `.docx` (one paragraph per line) before being
handed to the same `DocxJobProcessor` / `translate_docx` pipeline used for
native `.docx` uploads, and the resulting translated `.docx` is unwrapped back
to plain text for the final output. See
docs/architecture/pdf-vs-docx-translation-architecture.md for the rationale
behind reusing the DOCX pipeline for structurally simple formats.
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document as DocxDocument


def txt_bytes_to_docx_bytes(text_content: bytes) -> bytes:
    """Wrap plain-text bytes in a minimal `.docx` (one paragraph per line).

    Decoding is lenient (`errors="replace"`) so a `.txt` file with a few
    non-UTF-8 bytes still produces a translatable document instead of
    failing the whole submission.
    """
    text = text_content.decode("utf-8", errors="replace")
    document = DocxDocument()
    lines = text.splitlines() or [""]
    for line in lines:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_path_to_txt_bytes(docx_path: Path) -> bytes:
    """Extract plain text (one line per top-level paragraph) from a `.docx` file.

    Only top-level body paragraphs are read (matching what `add_paragraph`
    wrote in `txt_bytes_to_docx_bytes`), which also naturally includes any
    AI-translation cover page paragraphs prepended by the DOCX pipeline.
    """
    document = DocxDocument(str(docx_path))
    lines = [paragraph.text for paragraph in document.paragraphs]
    return "\n".join(lines).encode("utf-8")
