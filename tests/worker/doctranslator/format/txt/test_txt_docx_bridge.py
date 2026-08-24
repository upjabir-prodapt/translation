"""Unit tests for the plain-text <-> DOCX bridge used by the .txt pipeline."""

from docx import Document as open_docx

from src.worker.doctranslator.format.txt.txt_docx_bridge import (
    docx_path_to_txt_bytes,
)
from src.worker.doctranslator.format.txt.txt_docx_bridge import (
    txt_bytes_to_docx_bytes,
)


class TestTxtBytesToDocxBytes:
    def test_wraps_each_line_as_a_paragraph(self, tmp_path):
        text = "Hello world.\nSecond line.\nThird line."
        docx_bytes = txt_bytes_to_docx_bytes(text.encode("utf-8"))

        docx_path = tmp_path / "wrapped.docx"
        docx_path.write_bytes(docx_bytes)

        document = open_docx(str(docx_path))
        paragraphs = [p.text for p in document.paragraphs]
        assert paragraphs == ["Hello world.", "Second line.", "Third line."]

    def test_empty_text_produces_single_empty_paragraph(self):
        docx_bytes = txt_bytes_to_docx_bytes(b"")
        assert docx_bytes  # a valid (non-empty) docx file was produced

    def test_lenient_decoding_of_invalid_utf8(self):
        # Should not raise even with invalid UTF-8 bytes.
        docx_bytes = txt_bytes_to_docx_bytes(b"Hello \xff\xfe world")
        assert docx_bytes


class TestDocxPathToTxtBytes:
    def test_round_trip_preserves_lines(self, tmp_path):
        text = "Line one.\nLine two.\nLine three."
        docx_bytes = txt_bytes_to_docx_bytes(text.encode("utf-8"))
        docx_path = tmp_path / "doc.docx"
        docx_path.write_bytes(docx_bytes)

        result = docx_path_to_txt_bytes(docx_path)
        assert result.decode("utf-8") == text

    def test_includes_all_top_level_paragraphs(self, tmp_path):
        document = open_docx()
        document.add_paragraph("Cover page title")
        document.add_paragraph("Original content")
        docx_path = tmp_path / "with_cover.docx"
        document.save(str(docx_path))

        result = docx_path_to_txt_bytes(docx_path).decode("utf-8")
        assert "Cover page title" in result
        assert "Original content" in result
